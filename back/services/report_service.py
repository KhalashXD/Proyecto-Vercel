from datetime import datetime
from io import BytesIO
from typing import Any, Dict, Iterable, List

from sqlalchemy.orm import Session

from models import Incident
from services.incident_service import obtener_incidente_por_codigo


def _format_datetime(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y %H:%M")
    return str(value)


def _incident_rows(incidents: Iterable[Incident]) -> List[Dict[str, Any]]:
    rows = []
    for incident in incidents:
        rows.append(
            {
                "Codigo": incident.incident_code,
                "Clave": incident.emergency_type.code,
                "Tipo": incident.emergency_type.name,
                "Direccion": f"{incident.street_1} con {incident.street_2}",
                "Estado": incident.status,
                "Prioridad": incident.priority,
                "Creado": _format_datetime(incident.created_at),
                "Cerrado": _format_datetime(incident.closed_at),
                "Unidades": ", ".join(
                    relation.vehicle.vehicle_code for relation in incident.vehicles
                ),
            }
        )
    return rows


def obtener_filas_historial(db: Session) -> List[Dict[str, Any]]:
    incidents = (
        db.query(Incident)
        .order_by(Incident.created_at.desc())
        .all()
    )
    return _incident_rows(incidents)


def obtener_filas_incidente(db: Session, incident_code: str) -> List[Dict[str, Any]]:
    incident = obtener_incidente_por_codigo(db=db, incident_code=incident_code)
    rows = [
        {
            "Seccion": "Resumen",
            "Fecha": _format_datetime(incident["created_at"]),
            "Tipo": incident["emergency"]["code"],
            "Descripcion": incident["emergency"]["name"],
            "Unidad": "",
            "Usuario": "",
        }
    ]

    for vehicle in incident["assigned_vehicles"]:
        rows.append(
            {
                "Seccion": "Unidad asignada",
                "Fecha": _format_datetime(vehicle["assigned_at"]),
                "Tipo": vehicle["vehicle_type"],
                "Descripcion": vehicle["station_name"],
                "Unidad": vehicle["vehicle_code"],
                "Usuario": "",
            }
        )

    for event in incident["timeline"]:
        rows.append(
            {
                "Seccion": "Cronologia",
                "Fecha": _format_datetime(event["created_at"]),
                "Tipo": event["event_type"],
                "Descripcion": event["description"],
                "Unidad": "",
                "Usuario": event["user_name"] or "",
            }
        )

    return rows


def crear_excel(rows: List[Dict[str, Any]], sheet_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]

    headers = list(rows[0].keys()) if rows else ["Sin datos"]
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])

    for column in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        worksheet.column_dimensions[column[0].column_letter].width = min(
            max(max_length + 2, 12),
            45,
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def crear_pdf(rows: List[Dict[str, Any]], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )

    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    if rows:
        headers = list(rows[0].keys())
        data = [headers] + [
            [str(row.get(header, ""))[:120] for header in headers]
            for row in rows
        ]
    else:
        data = [["Sin datos"]]

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")]),
            ]
        )
    )
    elements.append(table)

    document.build(elements)
    return output.getvalue()
